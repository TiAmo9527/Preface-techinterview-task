"""IM-01 workbook boundary. No persistence, planning, or raw-file storage."""

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
import math
from pathlib import PurePath

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.services import validators as v
from src.services.contracts import Diagnostic


# Approved wire schemas, shared by the reader and blank templates.
ASSET_HEADERS = (
    "property_id", "property_name", "location", "city", "room_id", "room_number",
    "lighting_status", "lighting_observed_on", "lighting_recorder", "lighting_note",
    "water_supply_status", "water_supply_observed_on", "water_supply_recorder", "water_supply_note",
    "air_conditioning_status", "air_conditioning_observed_on", "air_conditioning_recorder",
    "air_conditioning_note", "asset_id", "asset_name", "facility_type", "purchase_date",
    "installation_date", "useful_life_months", "acquisition_cost", "currency",
)
INVOICE_HEADERS = (
    "invoice_id", "line_id", "room_id", "facility_type", "asset_name", "purchase_date",
    "installation_date", "useful_life_months", "acquisition_cost", "currency",
    "invoice_date", "supplier_name",
)


def headers_for(workflow: str) -> tuple[str, ...]:
    if workflow == "ASSETS":
        return ASSET_HEADERS
    if workflow == "INVOICES":
        return INVOICE_HEADERS
    raise ValueError("Choose ASSETS or INVOICES")


Scalar = str | int | float | bool | None


@dataclass(frozen=True)
class SourceCell:
    field: str
    coordinate: str
    value: Scalar
    data_type: str


@dataclass(frozen=True)
class SourceRow:
    file: str
    sheet: str
    row: int
    cells: tuple[SourceCell, ...]


@dataclass(frozen=True)
class NormalizedRow:
    source: SourceRow
    values: tuple[tuple[str, str | int | None], ...]

    def as_dict(self) -> dict[str, str | int | None]:
        """A detached working copy; callers cannot mutate retained input."""
        return dict(self.values)


@dataclass(frozen=True)
class ParsedWorkbook:
    workflow: str
    filename: str
    source_rows: tuple[SourceRow, ...]
    normalized_rows: tuple[NormalizedRow, ...]
    diagnostics: tuple[Diagnostic, ...]

    @property
    def blockers(self) -> tuple[Diagnostic, ...]:
        return tuple(item for item in self.diagnostics if item.severity == "BLOCKER")

    @property
    def warnings(self) -> tuple[Diagnostic, ...]:
        return tuple(item for item in self.diagnostics if item.severity == "WARNING")


def _display(value: object) -> Scalar:
    """Keep offending information JSON-safe and deeply immutable."""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if type(value) in (str, int, bool, type(None)):
        return value
    if type(value) is float and math.isfinite(value):
        return value
    return repr(value)


def _empty(value: object) -> bool:
    return value is None or (type(value) is str and not value.strip())


def _normalize(raw: dict, source: SourceRow, workflow: str,
               diagnostics: list[Diagnostic]) -> NormalizedRow | None:
    values = {}
    invalid = set()
    identity = raw.get("asset_id") if workflow == "ASSETS" else (
        f"{raw.get('invoice_id')}/{raw.get('line_id')}"
    )

    def issue(field: str, reason: str, *, severity: str = "BLOCKER") -> None:
        diagnostics.append(Diagnostic(
            severity=severity, entity="asset" if workflow == "ASSETS" else "invoice_item",
            file=source.file, sheet=source.sheet, row=source.row, field=field,
            identity=None if identity is None else str(identity),
            value=_display(raw.get(field)), reason=reason,
        ))
        if severity == "BLOCKER":
            invalid.add(field)

    def apply(field: str, function, **kwargs) -> None:
        try:
            values[field] = function(raw[field], **kwargs)
        except (ValueError, TypeError) as error:
            issue(field, str(error))

    for cell in source.cells:
        if cell.data_type == "e":
            issue(cell.field, "Replace Excel error cells with valid source values")
    for field in raw:
        if field in invalid:
            continue
        if field in ("property_id", "room_id", "asset_id"):
            apply(field, v.source_identity, kind=field.removesuffix("_id"))
        elif field in ("purchase_date", "installation_date", "invoice_date") or field.endswith("_observed_on"):
            apply(field, v.calendar_date, optional=field == "installation_date" or field.endswith("_observed_on"))
        elif field == "acquisition_cost":
            apply(field, v.decimal_amount)
        elif field == "useful_life_months":
            apply(field, v.whole_months)
        elif field in ("currency", "location", "facility_type"):
            choices = {"currency": v.CURRENCIES, "location": v.LOCATIONS, "facility_type": v.SYSTEMS}
            apply(field, v.enum_value, choices=choices[field])
        elif field.endswith("_status"):
            if _empty(raw[field]):
                values[field] = None
            else:
                apply(field, v.enum_value, choices=v.CONDITIONS)
        else:
            apply(field, v.text, optional=field == "supplier_name" or field.endswith(("_recorder", "_note")))

    if "purchase_date" not in invalid and "installation_date" not in invalid:
        try:
            v.date_order(values["purchase_date"], values["installation_date"])
        except ValueError as error:
            issue("installation_date", str(error))

    if workflow == "ASSETS":
        for field, kwargs in (
            ("property_id", {"location": "location"}),
            ("room_id", {"parent": "property_id"}),
            ("asset_id", {"parent": "room_id"}),
        ):
            if not invalid.intersection((field, *kwargs.values())):
                try:
                    v.source_identity(values[field], field.removesuffix("_id"),
                                      **{key: values[name] for key, name in kwargs.items()})
                except ValueError as error:
                    issue(field, str(error))
        for system in ("lighting", "water_supply", "air_conditioning"):
            fields = tuple(f"{system}_{suffix}" for suffix in ("status", "observed_on", "recorder", "note"))
            if invalid.intersection(fields):
                continue
            try:
                group = v.observation_metadata(*(values[field] for field in fields))
                for field, suffix in zip(fields, ("status", "observed_on", "recorder", "note")):
                    values[field] = group[suffix]
            except ValueError as error:
                # Point at every missing required group field, rather than hiding it in one row error.
                for field in fields[:3]:
                    if values[field] is None:
                        issue(field, str(error))

    if invalid:
        return None
    if values["installation_date"] is None:
        issue("installation_date", "Installation is absent; purchase date supplies the service fallback", severity="WARNING")
    if values["acquisition_cost"] == "0":
        issue("acquisition_cost", "Explicit zero cost is retained", severity="WARNING")
    if workflow == "INVOICES" and values["supplier_name"] is None:
        issue("supplier_name", "Supplier is absent; provide one when known", severity="WARNING")
    return NormalizedRow(source, tuple(values.items()))


def parse_workbook(content: bytes, *, filename: str, workflow: str) -> ParsedWorkbook:
    """Read once, retain all nonempty source rows, and never access a store."""
    diagnostics = []
    source_rows = []
    normalized_rows = []

    def issue(field: str, reason: str, *, sheet: str | None = None,
              row: int | None = None, value: object = None) -> None:
        diagnostics.append(Diagnostic(entity="workbook", file=filename or None,
                                      sheet=sheet, row=row, field=field,
                                      value=_display(value), reason=reason))

    def result() -> ParsedWorkbook:
        return ParsedWorkbook(workflow, filename, tuple(source_rows),
                              tuple(normalized_rows), tuple(diagnostics))

    try:
        headers = headers_for(workflow)
    except ValueError as error:
        issue("workflow", str(error), value=workflow)
        return result()
    if not filename or PurePath(filename).suffix.lower() != ".xlsx":
        issue("file", "Select an .xlsx workbook for this workflow", value=filename)
        return result()
    if not content:
        issue("file", "The workbook is missing or empty")
        return result()
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    except Exception:
        # Workbook decoding errors vary by ZIP/XML/openpyxl layer. Never expose internals.
        issue("file", "The workbook cannot be read; supply a readable .xlsx file")
        return result()
    try:
        expected_sheet = "Assets" if workflow == "ASSETS" else "Invoices"
        if workbook.sheetnames != [expected_sheet]:
            issue("sheet", f"Require exactly one sheet named {expected_sheet}", value=repr(workbook.sheetnames))
            return result()
        sheet = workbook[expected_sheet]
        rows = tuple(sheet.iter_rows())
        if not rows:
            rows = (tuple(sheet.cell(1, column) for column in range(1, len(headers) + 1)),)
        # Styles alone do not create source fields/rows; nonempty values do.
        width = max((cell.column for row in rows for cell in row if cell.value is not None), default=len(headers))
        header_cells = rows[0][:width]
        positions = {}
        for cell in header_cells:
            value = cell.value
            if type(value) is not str or value not in headers:
                issue(get_column_letter(cell.column), "Unexpected or blank header; use the exact template headers",
                      sheet=expected_sheet, row=1, value=value)
            elif value in positions:
                issue(value, "Duplicate header; include each header exactly once", sheet=expected_sheet, row=1, value=value)
            else:
                positions[value] = cell.column - 1
        for field in headers:
            if field not in positions:
                issue(field, "Missing required header; include optional-value columns too", sheet=expected_sheet, row=1)
        for merged in sheet.merged_cells.ranges:
            issue("merged_cells", "Unmerge table cells before importing", sheet=expected_sheet,
                  row=merged.min_row, value=str(merged))
        for row in rows:
            for cell in row:
                if cell.data_type == "f":
                    field = next((key for key, index in positions.items() if index == cell.column - 1), get_column_letter(cell.column))
                    issue(field, "Replace formulas with source values", sheet=expected_sheet, row=cell.row, value=cell.value)
        structure_valid = not diagnostics
        for row in rows[1:]:
            if all(_empty(cell.value) for cell in row):
                continue
            source = SourceRow(filename, expected_sheet, row[0].row, tuple(
                SourceCell(next((key for key, index in positions.items() if index == cell.column - 1), get_column_letter(cell.column)),
                           cell.coordinate, _display(cell.value), cell.data_type)
                for cell in row[:width]
            ))
            source_rows.append(source)
            if structure_valid:
                raw = {field: row[positions[field]].value for field in headers}
                normalized = _normalize(raw, source, workflow, diagnostics)
                if normalized is not None:
                    normalized_rows.append(normalized)
        return result()
    finally:
        workbook.close()
